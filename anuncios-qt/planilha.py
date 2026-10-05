#!/usr/bin/env python3
"""
Gera anuncios-qt/meta-ads-qt.xlsx, a planilha de Meta Ads da QT Pizza Bar.

Seis abas:
  Calculadora  quanto a QT pode pagar por cliente, a projecao do mes e o teste de estresse
  Campanhas    a ficha tecnica de cada campanha
  Semanal      o acompanhamento de toda segunda
  Criativos    qual gancho e qual angulo funcionam
  Checklist    semana 0 e lancamento
  Como ler     glossario e regras de decisao

Identidade da QT: preto #1A1E1E, cinza #A0A5A5, branco #EFECEC, Arial no lugar
da Helvetica. Amarelo e o que se preenche, branco QT e o que se calcula.

Os numeros que vem preenchidos sao os mesmos exemplos da apostila. Troque pelos
da QT direto na planilha, ou aqui e rode de novo.

Uso: python3 anuncios-qt/planilha.py
"""
import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

PRETO="FF1A1E1E"; CINZA="FFA0A5A5"; BRANCO="FFEFECEC"; AZUL="FF0000FF"
F="Arial"
PREENCHER=PatternFill("solid", fgColor="FFFFF2CC")   # amarelo: voce preenche
CALC=PatternFill("solid", fgColor=BRANCO)             # branco QT: calculado
CAB=PatternFill("solid", fgColor=PRETO)
VERDE_CLARO=PatternFill("solid", fgColor="FFDDEBDD")
AMARELO_CLARO=PatternFill("solid", fgColor="FFF6EBC8")
VERMELHO_CLARO=PatternFill("solid", fgColor="FFF2D4D0")
fina=Side(style="thin", color="FFD5D3D0")
BORDA=Border(left=fina,right=fina,top=fina,bottom=fina)

R='"R$" #,##0.00;[Red]-"R$" #,##0.00;"-"'
PCT="0.0%"; PCT2="0.00%"; INT="#,##0"; DEC1="0.0"; DEC2="0.00"; FATOR="0.0000"
DATA="dd/mm/yyyy"

wb=Workbook()

def nome(n, ref):
    wb.defined_names[n]=DefinedName(n, attr_text=ref)

def titulo(ws, texto, sub):
    ws["A1"]=texto; ws["A1"].font=Font(F,size=14,bold=True,color=PRETO)
    ws["A2"]=sub; ws["A2"].font=Font(F,size=10,italic=True,color=CINZA)

def cabecalho(ws, linha, nomes, col=1, altura=32):
    for i,n in enumerate(nomes,col):
        c=ws.cell(linha,i,n); c.font=Font(F,size=10,bold=True,color="FFFFFFFF")
        c.fill=CAB; c.border=BORDA
        c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    ws.row_dimensions[linha].height=altura

def secao(ws, linha, texto):
    c=ws.cell(linha,1,texto); c.font=Font(F,size=11,bold=True,color=PRETO)

def rotulo(ws, linha, texto, negrito=False):
    c=ws.cell(linha,1,texto); c.font=Font(F,size=10,bold=negrito)
    c.alignment=Alignment(vertical="center",wrap_text=True); c.border=BORDA

def entrada(c, valor, fmt=None):
    c.value=valor; c.fill=PREENCHER; c.font=Font(F,size=10,color=AZUL); c.border=BORDA
    c.alignment=Alignment(vertical="center")
    if fmt: c.number_format=fmt

def calculo(c, formula, fmt=None, negrito=False):
    c.value=formula; c.fill=CALC; c.font=Font(F,size=10,bold=negrito); c.border=BORDA
    c.alignment=Alignment(vertical="center")
    if fmt: c.number_format=fmt

def nota(ws, linha, texto, de=3, ate=6):
    ws.merge_cells(start_row=linha,start_column=de,end_row=linha,end_column=ate)
    c=ws.cell(linha,de,texto); c.font=Font(F,size=9,italic=True,color=CINZA)
    c.alignment=Alignment(vertical="center",wrap_text=True)
    if len(texto)>70: ws.row_dimensions[linha].height=26

def exemplo(c, valor, fmt=None):
    c.value=valor; c.font=Font(F,size=9,italic=True,color=CINZA); c.border=BORDA
    if fmt: c.number_format=fmt


# ---------- calculadora ----------
ws=wb.active; ws.title="Calculadora"
titulo(ws,"Calculadora · quanto a QT pode pagar por um cliente",
       "Preencha só o amarelo. Os números que vêm preenchidos são exemplo, os mesmos da apostila: troque pelos da QT. O resto se calcula.")

secao(ws,4,"1 · A casa")
linhas_casa=[
 (5,"Ticket médio por pessoa, sem os 13%",150,R,"Exemplo. No Altec: faturamento sem serviço ÷ número de clientes. Os 13% são da equipe e não entram."),
 (6,"CMV médio, comida e bebida",0.28,PCT,"Exemplo. CMV médio do mês, somando comida e bebida."),
 (7,"Impostos sobre a venda",0.10,PCT,"Exemplo. Depende do regime tributário da QT: confirme com o contador."),
 (8,"Cartão e outros custos variáveis",0.03,PCT,"Exemplo. Taxa de cartão e o que mais varia com cada venda."),
]
for r,t,v,fmt,n in linhas_casa:
    rotulo(ws,r,t); entrada(ws.cell(r,2),v,fmt); nota(ws,r,n)
rotulo(ws,9,"Margem de contribuição",True); calculo(ws["B9"],"=1-B6-B7-B8",PCT,True)
nota(ws,9,"O que sobra de cada real vendido para pagar o fixo e, agora, o anúncio.")
rotulo(ws,10,"Contribuição por pessoa",True); calculo(ws["B10"],"=B5*B9",R,True)
nota(ws,10,"Quanto cada pessoa sentada deixa para a casa.")

secao(ws,12,"2 · O funil depois do anúncio")
linhas_funil=[
 (13,"Conversa que vira reserva",0.25,PCT,"Etiqueta Reservou ÷ conversas. Comece com 25% e troque pelo real na semana 3."),
 (14,"Comparecimento",0.85,PCT,"Mesas que vieram ÷ reservas feitas. 85% é o mesmo que 15% de não comparecimento."),
 (15,"Pessoas por reserva",2.5,DEC1,"Média de pessoas por mesa reservada."),
]
for r,t,v,fmt,n in linhas_funil:
    rotulo(ws,r,t); entrada(ws.cell(r,2),v,fmt); nota(ws,r,n)
rotulo(ws,16,"Contribuição por mesa que veio"); calculo(ws["B16"],"=B10*B15",R)
nota(ws,16,"Contribuição por pessoa × pessoas por reserva.")
rotulo(ws,17,"Contribuição por reserva feita"); calculo(ws["B17"],"=B16*B14",R)
nota(ws,17,"Já descontado quem reservou e não veio.")
rotulo(ws,18,"Contribuição por conversa",True); calculo(ws["B18"],"=B17*B13",R,True)
nota(ws,18,"Se a conversa custar isso, com tributo, a QT empata na primeira visita.")

secao(ws,20,"3 · O tributo da Meta e a margem de segurança")
rotulo(ws,21,"Tributo repassado pela Meta desde 01/01/2026"); entrada(ws["B21"],0.1215,PCT2)
nota(ws,21,"PIS e Cofins 9,25% mais ISS 2,9%. Cobrado por dentro.")
rotulo(ws,22,"Fator do tributo"); calculo(ws["B22"],"=1/(1-B21)",FATOR)
nota(ws,22,"Valor do Gerenciador × fator = fatura. R$ 100 vira R$ 113,83.")
rotulo(ws,23,"Meta, em fração do empate"); entrada(ws["B23"],0.5,PCT)
nota(ws,23,"Pagar no máximo metade do empate cobre cliente que viria de qualquer jeito e erro de medição. Apostila, módulo 7.3.")

secao(ws,25,"4 · Quanto dá para pagar")
cabecalho(ws,26,["Por","Empate real, com tributo","Meta real","Empate no Gerenciador","Meta no Gerenciador"])
for r,t,origem in [(27,"Conversa","B18"),(28,"Reserva feita","B17"),(29,"Pessoa atendida","B10")]:
    rotulo(ws,r,t,True)
    calculo(ws.cell(r,2),f"={origem}",R)
    calculo(ws.cell(r,3),f"=B{r}*$B$23",R,r==29)
    calculo(ws.cell(r,4),f"=B{r}/$B$22",R)
    calculo(ws.cell(r,5),f"=D{r}*$B$23",R,r==27)
nota(ws,30,"No Gerenciador é o real ÷ fator do tributo, para comparar direto com o que a tela mostra. Os dois em negrito são os da semana: custo por conversa do Gerenciador contra a meta da linha Conversa, custo real por pessoa contra a meta real da linha Pessoa.",1,6)
ws.row_dimensions[30].height=26
rotulo(ws,31,"ROAS real de empate"); calculo(ws["B31"],"=1/B9",DEC2)
nota(ws,31,"Cada real de anúncio precisa trazer isso de faturamento para empatar.")
rotulo(ws,32,"ROAS real meta"); calculo(ws["B32"],"=1/(B9*B23)",DEC2,True)
nota(ws,32,"O ROAS que a QT quer ver na aba Semanal.")
rotulo(ws,33,"% de mídia de empate"); calculo(ws["B33"],"=B9",PCT)
nota(ws,33,"O CMV do marketing: gasto real ÷ faturamento que veio do anúncio.")
rotulo(ws,34,"% de mídia meta"); calculo(ws["B34"],"=B9*B23",PCT,True)

secao(ws,36,"5 · Projeção do mês")
linhas_proj=[
 (37,"Verba diária da campanha de reserva, no Gerenciador",50,R,"Campanha A, QT_RES_WPP."),
 (38,"Verba diária da vitrine, no Gerenciador",15,R,"Campanha B, QT_VIT_ALC. Não gera conversa, mas entra no custo."),
 (39,"Dias no mês",30,INT,""),
 (40,"Custo por conversa esperado, no Gerenciador",12,R,"Hipótese prudente. O mercado fala em perto de R$ 5 no Brasil em 2026. Troque pelo real da QT na semana 3."),
]
for r,t,v,fmt,n in linhas_proj:
    rotulo(ws,r,t); entrada(ws.cell(r,2),v,fmt)
    if n: nota(ws,r,n)
proj=[
 (41,"Verba do mês no Gerenciador","=(B37+B38)*B39",R,False,""),
 (42,"Desembolso real do mês, com tributo","=B41*B22",R,True,"O que sai do caixa."),
 (43,"Conversas","=B37*B39/B40",INT,False,"Só a campanha de reserva gera conversa."),
 (44,"Reservas","=B43*B13",INT,False,""),
 (45,"Mesas que vieram","=B44*B14",INT,False,""),
 (46,"Pessoas atendidas","=B45*B15",INT,False,""),
 (47,"Faturamento","=B46*B5",R,False,"Pessoas × ticket médio."),
 (48,"Contribuição","=B47*B9",R,False,""),
 (49,"Resultado do mês","=B48-B42",R,True,"Contribuição menos o que saiu do caixa."),
 (50,"ROAS real","=B47/B42",DEC2,False,""),
 (51,"% de mídia","=B42/B47",PCT,False,""),
 (52,"Custo real por pessoa atendida","=B42/B46",R,True,"A métrica que manda. Compare com a linha Pessoa atendida da seção 4."),
 (53,"Leitura",'=IF(B52<=C29,"dentro da meta",IF(B52<=B29,"entre a meta e o empate","acima do empate: a conta não fecha"))',None,True,""),
]
for r,t,f_,fmt,neg,n in proj:
    rotulo(ws,r,t,neg); calculo(ws.cell(r,2),f_,fmt,neg)
    if n: nota(ws,r,n)
rotulo(ws,54,"Lugares que sobram por semana nos dias fracos"); entrada(ws["B54"],80,INT)
nota(ws,54,"Exemplo. Quantas pessoas a mais cabem de terça a quinta, somando os três dias.")
rotulo(ws,55,"Pessoas do anúncio por semana"); calculo(ws["B55"],"=B46/(B39/7)",DEC1)
rotulo(ws,56,"Parte da sobra que o anúncio ocupa"); calculo(ws["B56"],"=B55/B54",PCT)
nota(ws,56,"Acima de 100%, o anúncio traz mais gente do que cabe nos dias fracos.")

secao(ws,58,"6 · Teste de estresse · resultado do mês com a verba da seção 5")
c=ws.cell(59,1,"Custo por conversa no Gerenciador, para baixo · conversa que vira reserva, para o lado")
c.font=Font(F,size=9,bold=True,color="FFFFFFFF"); c.fill=CAB; c.border=BORDA
c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
ws.row_dimensions[59].height=32
taxas=[0.15,0.20,0.25,0.30,0.35]
custos=[6,8,10,12,15,20,25,30]
for j,t in enumerate(taxas,2): entrada(ws.cell(59,j),t,PCT)
for i,cpc in enumerate(custos,60):
    entrada(ws.cell(i,1),cpc,R)
    for j in range(2,2+len(taxas)):
        col=get_column_letter(j)
        calculo(ws.cell(i,j),f"=($B$37*$B$39/$A{i})*{col}$59*$B$14*$B$15*$B$5*$B$9-$B$42",R)
fim_grade=59+len(custos)
ws.conditional_formatting.add(f"B60:F{fim_grade}",CellIsRule(operator="lessThan",formula=["0"],fill=VERMELHO_CLARO))
nota(ws,fim_grade+1,"Vermelho é prejuízo. Repare que dobrar a conversão em reserva vale o mesmo que cortar o custo por conversa pela metade, e a conversão depende do atendimento.",1,6)
ws.row_dimensions[fim_grade+1].height=26

for col,w in zip("ABCDEF",[48,17,17,17,17,17]): ws.column_dimensions[col].width=w
ws.sheet_properties.tabColor=PRETO

nome("ticket","Calculadora!$B$5")
nome("margem","Calculadora!$B$9")
nome("fator_tributo","Calculadora!$B$22")
nome("meta_conversa","Calculadora!$E$27")
nome("meta_pessoa","Calculadora!$C$29")
nome("empate_pessoa","Calculadora!$B$29")
nome("verba_reserva","Calculadora!$B$37")
nome("verba_vitrine","Calculadora!$B$38")


# ---------- campanhas ----------
s2=wb.create_sheet("Campanhas")
titulo(s2,"Campanhas · a ficha técnica de cada uma",
       "O detalhe de cada ajuste está no plano de largada. A verba em amarelo é a que você vai colocar no Gerenciador.")
cab2=["Campanha","Objetivo","Para onde a pessoa vai","Meta de desempenho","Conjunto","Local e raio",
      "Idade","Público","Verba diária no Gerenciador","Anúncios","Mensagem pré-preenchida","Status"]
cabecalho(s2,4,cab2)
camps=[
 ("QT_RES_WPP","Engajamento","WhatsApp da QT","Maximizar o número de conversas","R6km_18+_ADV",
  "endereço da QT, 6 km","mínima 18, sugestão 25 a 45","Advantage+",50,"A1, A2, A3, A4",
  "Oi! Vi o anúncio da QT e quero reservar uma mesa.","planejada"),
 ("QT_VIT_ALC","Reconhecimento","botão Saiba mais com link wa.me","Maximizar o alcance, no máximo 2 impressões a cada 7 dias",
  "R4km_18+_ADV","endereço da QT, 4 km","mínima 18, sugestão 25 a 45","Advantage+",15,"B1, B2",
  "Oi! Vi o vídeo da QT e quero reservar.","planejada"),
 ("QT_GRP_WPP","Engajamento","WhatsApp da QT","Maximizar o número de conversas","R5km_18+_ADV",
  "endereço da QT, 3 a 6 km","mínima 18","Advantage+",0,"a criar, liga no fim de outubro",
  "Oi! Quero orçar uma confraternização na QT.","sazonal"),
]
for n,row in enumerate(camps):
    r=5+n
    for i,v in enumerate(row,1):
        c=s2.cell(r,i)
        if i==9:
            entrada(c,v,R)
        elif i==12:
            entrada(c,v)
        else:
            c.value=v; c.font=Font(F,size=10,bold=(i==1)); c.border=BORDA
        c.alignment=Alignment(vertical="center",wrap_text=True)
    s2.row_dimensions[r].height=46
rotulo(s2,8,"Total por dia",True)
calculo(s2["I8"],"=SUM(I5:I7)",R,True)
c=s2["K8"]
c.value='=IF(AND(I5=verba_reserva,I6=verba_vitrine),"reserva e vitrine batem com a Calculadora","diferente da Calculadora: atualize lá para a projeção valer")'
c.font=Font(F,size=9,italic=True,color=CINZA); c.alignment=Alignment(wrap_text=True,vertical="center")
s2.merge_cells("K8:L8"); s2.row_dimensions[8].height=30
calculo(s2["I9"],"=I8*fator_tributo*30",R)
rotulo(s2,9,"Fatura de 30 dias, com tributo")

dv_status=DataValidation(type="list",formula1='"planejada,ativa,pausada,encerrada,sazonal"',allow_blank=True)
s2.add_data_validation(dv_status); dv_status.add("L5:L7")

secao(s2,11,"Nomenclatura")
cabecalho(s2,12,["Nível","Molde","Exemplo","Lê-se"])
nomencl=[
 ("Campanha","QT_[objetivo]_[destino]","QT_RES_WPP","reserva pelo WhatsApp"),
 ("Conjunto","[local]_[idade]_[público]","R6km_18+_ADV","raio de 6 km, 18+, Advantage+"),
 ("Anúncio","[ângulo]_[formato]_[gancho]_[versão]","DEST_REELS_FORNO_v1","destino, Reels, abre no forno, versão 1"),
 ("Ângulos","DEST · OCAS · PROC · PROD · VIT","","destino, ocasião, processo, produto, vitrine"),
]
for n,row in enumerate(nomencl):
    for i,v in enumerate(row,1):
        c=s2.cell(13+n,i,v); c.font=Font(F,size=10,bold=(i==1)); c.border=BORDA
        c.alignment=Alignment(vertical="center",wrap_text=True)
    s2.row_dimensions[13+n].height=30
for col,w in zip("ABCDEFGHIJKL",[16,16,24,28,16,20,18,13,15,14,34,14]):
    s2.column_dimensions[col].width=w


# ---------- semanal ----------
s3=wb.create_sheet("Semanal")
titulo(s3,"Acompanhamento semanal · toda segunda, os últimos 7 dias",
       "Amarelo: copie do Gerenciador (predefinição QT semanal), das etiquetas do WhatsApp e do livro de reservas. Branco QT: calcula sozinho. A linha EXEMPLO é só formato: apague quando começar.")
s3["A4"]="Segunda da semana 1:"; s3["A4"].font=Font(F,size=10,bold=True)
entrada(s3["D4"],datetime.date(2026,10,12),DATA)
s3["E4"]="troque pela segunda da semana do lançamento. As outras datas se ajustam."
s3["E4"].font=Font(F,size=9,italic=True,color=CINZA)

cab3=["#","Semana de",
      "Gasto total no Gerenciador","Gasto da reserva no Gerenciador","Impressões da reserva","Cliques no link da reserva",
      "Conversas iniciadas","Alcance da vitrine","Frequência da vitrine",
      "Reservas feitas","Mesas que vieram","Pessoas atendidas",
      "Gasto real, com tributo","CPM da reserva","CTR do link","Custo por conversa no Gerenciador",
      "Conversa que vira reserva","Comparecimento","Custo real por pessoa",
      "Faturamento estimado","Contribuição","Resultado da semana","ROAS real","% de mídia",
      "Leitura","Mudança feita nesta segunda"]
L3=6
cabecalho(s3,L3,cab3,altura=48)
ENTRADAS3=list(range(3,13))   # C a L

def formulas_semana(r):
    return {
     13:(f'=IF(C{r}="","",C{r}*fator_tributo)',R),
     14:(f'=IFERROR(IF(OR(D{r}="",E{r}=""),"",D{r}/E{r}*1000),"")',R),
     15:(f'=IFERROR(IF(OR(F{r}="",E{r}=""),"",F{r}/E{r}),"")',PCT2),
     16:(f'=IFERROR(IF(OR(D{r}="",G{r}=""),"",D{r}/G{r}),"")',R),
     17:(f'=IFERROR(IF(OR(J{r}="",G{r}=""),"",J{r}/G{r}),"")',PCT),
     18:(f'=IFERROR(IF(OR(K{r}="",J{r}=""),"",K{r}/J{r}),"")',PCT),
     19:(f'=IFERROR(IF(OR(M{r}="",L{r}=""),"",M{r}/L{r}),"")',R),
     20:(f'=IF(L{r}="","",L{r}*ticket)',R),
     21:(f'=IF(T{r}="","",T{r}*margem)',R),
     22:(f'=IF(OR(U{r}="",M{r}=""),"",U{r}-M{r})',R),
     23:(f'=IFERROR(IF(OR(T{r}="",M{r}=""),"",T{r}/M{r}),"")',DEC2),
     24:(f'=IFERROR(IF(OR(M{r}="",T{r}=""),"",M{r}/T{r}),"")',PCT),
     25:(f'=IF(S{r}="","",IF(S{r}<=meta_pessoa,"dentro da meta",IF(S{r}<=empate_pessoa,"entre a meta e o empate","acima do empate")))',None),
    }

FMT_ENTRADA3={3:R,4:R,5:INT,6:INT,7:INT,8:INT,9:DEC2,10:INT,11:INT,12:INT}

# linha de exemplo, cinza e italico
r=L3+1
exemplo(s3.cell(r,1),"ex"); exemplo(s3.cell(r,2),"EXEMPLO")
for i,v in zip(ENTRADAS3,[455,350,25000,210,29,6200,1.9,7,6,15]):
    exemplo(s3.cell(r,i),v,FMT_ENTRADA3[i])
for i,(f_,fmt) in formulas_semana(r).items():
    c=s3.cell(r,i,f_); c.font=Font(F,size=9,italic=True,color=CINZA); c.fill=CALC; c.border=BORDA
    if fmt: c.number_format=fmt
exemplo(s3.cell(r,26),"nenhuma, semana de aprendizado")

SEMANAS=12
ini3=L3+2
for n in range(SEMANAS):
    r=ini3+n
    s3.cell(r,1,n+1).font=Font(F,size=10); s3.cell(r,1).border=BORDA
    d=s3.cell(r,2,"=$D$4" if n==0 else f"=B{r-1}+7")
    d.font=Font(F,size=10); d.number_format="dd/mm"; d.border=BORDA
    for i in ENTRADAS3: entrada(s3.cell(r,i),None,FMT_ENTRADA3[i])
    for i,(f_,fmt) in formulas_semana(r).items(): calculo(s3.cell(r,i),f_,fmt)
    entrada(s3.cell(r,26),None)
fim3=ini3+SEMANAS-1

rt=fim3+2
c=s3.cell(rt,2,"TOTAL"); c.font=Font(F,size=10,bold=True)
for i in [3,4,5,6,7,10,11,12]:
    col=get_column_letter(i)
    calculo(s3.cell(rt,i),f"=SUM({col}{ini3}:{col}{fim3})",FMT_ENTRADA3[i],True)
for i in [13,20,21,22]:
    col=get_column_letter(i)
    calculo(s3.cell(rt,i),f"=SUM({col}{ini3}:{col}{fim3})",R,True)
tot={
 14:(f'=IFERROR(D{rt}/E{rt}*1000,"")',R),
 15:(f'=IFERROR(F{rt}/E{rt},"")',PCT2),
 16:(f'=IFERROR(D{rt}/G{rt},"")',R),
 17:(f'=IFERROR(J{rt}/G{rt},"")',PCT),
 18:(f'=IFERROR(K{rt}/J{rt},"")',PCT),
 19:(f'=IFERROR(M{rt}/L{rt},"")',R),
 23:(f'=IFERROR(T{rt}/M{rt},"")',DEC2),
 24:(f'=IFERROR(M{rt}/T{rt},"")',PCT),
 25:(f'=IF(S{rt}="","",IF(S{rt}<=meta_pessoa,"dentro da meta",IF(S{rt}<=empate_pessoa,"entre a meta e o empate","acima do empate")))',None),
}
for i,(f_,fmt) in tot.items(): calculo(s3.cell(rt,i),f_,fmt,True)
c=s3.cell(rt+1,2,"As taxas do total saem das somas, não da média das semanas.")
c.font=Font(F,size=9,italic=True,color=CINZA)

for txt,fill in [("dentro da meta",VERDE_CLARO),("entre a meta e o empate",AMARELO_CLARO),("acima do empate",VERMELHO_CLARO)]:
    s3.conditional_formatting.add(f"Y{ini3}:Y{rt}",CellIsRule(operator="equal",formula=[f'"{txt}"'],fill=fill))

larg3=[5,11]+[13]*10+[13]*12+[22,34]
for i,w in enumerate(larg3,1): s3.column_dimensions[get_column_letter(i)].width=w
s3.freeze_panes=f"C{L3+1}"
s3.sheet_properties.tabColor=PRETO


# ---------- criativos ----------
s4=wb.create_sheet("Criativos")
titulo(s4,"Criativos · qual gancho e qual ângulo funcionam",
       "Uma linha por anúncio, com os números acumulados desde que entrou no ar: Gerenciador, nível de anúncio, período máximo. Atualize quando for decidir se pausa e na revisão do mês.")
cab4=["Código","Campanha","Ângulo","Formato","Gancho, os 3 primeiros segundos","No ar desde",
      "Gasto no Gerenciador","Impressões","Reproduções de 3 s","ThruPlays","Cliques no link","Conversas",
      "Taxa de gancho","Retenção","CTR do link","Custo por conversa","Sugestão"]
L4=4
cabecalho(s4,L4,cab4,altura=40)

def formulas_criativo(r):
    sug=(f'=IF(H{r}="","",IF(H{r}<1000,"cedo para julgar",'
         f'IF(LEFT(B{r},6)="QT_VIT",IF(AND(I{r}<>"",M{r}<0.2),"refazer os 3 primeiros segundos","vitrine: julgar por alcance e gancho"),'
         f'IF(L{r}=0,IF(G{r}>=2*meta_conversa,"pausar: gastou 2x a meta sem conversa","ainda sem conversa: observar"),'
         f'IF(AND(I{r}<>"",M{r}<0.2),"refazer os 3 primeiros segundos",'
         f'IF(AND(K{r}<>"",O{r}<0.005),"rever texto, oferta e botão",'
         f'IF(P{r}<=meta_conversa,"manter: dentro da meta","observar: acima da meta")))))))')
    return {
     13:(f'=IFERROR(IF(OR(H{r}="",I{r}=""),"",I{r}/H{r}),"")',PCT),
     14:(f'=IFERROR(IF(OR(I{r}="",J{r}=""),"",J{r}/I{r}),"")',PCT),
     15:(f'=IFERROR(IF(OR(H{r}="",K{r}=""),"",K{r}/H{r}),"")',PCT2),
     16:(f'=IFERROR(IF(OR(G{r}="",L{r}="",L{r}=0),"",G{r}/L{r}),"")',R),
     17:(sug,None),
    }

FMT_ENTRADA4={7:R,8:INT,9:INT,10:INT,11:INT,12:INT}
r=L4+1
for i,v in enumerate(["EXEMPLO_REELS_v1","QT_RES_WPP","Destino","Reels 9:16","exemplo, apague",None,420,30000,9300,2100,260,38],1):
    exemplo(s4.cell(r,i),v,FMT_ENTRADA4.get(i))
for i,(f_,fmt) in formulas_criativo(r).items():
    c=s4.cell(r,i,f_); c.font=Font(F,size=9,italic=True,color=CINZA); c.fill=CALC; c.border=BORDA
    if fmt: c.number_format=fmt

planejados=[
 ("DEST_REELS_FORNO_v1","QT_RES_WPP","Destino","Reels 9:16","pizza entra no forno, borda estufando, 400°C na tela"),
 ("OCAS_REELS_SALAO_v1","QT_RES_WPP","Ocasião","Reels 9:16","salão às 20h, câmera atravessando até a mesa"),
 ("PROC_CARROSSEL_48H_v1","QT_RES_WPP","Processo","Carrossel 4:5","disco de massa aberto com as bolhas, 48 horas"),
 ("DEST_FOTO_MARGHERITA_v1","QT_RES_WPP","Destino","Foto 4:5","Margherita vista de cima, selo do 50 Top Pizza"),
 ("VIT_REELS_BORDA_v1","QT_VIT_ALC","Vitrine","Reels 9:16","borda estufando em tempo real, 400°C"),
 ("VIT_REELS_SALAO_v1","QT_VIT_ALC","Vitrine","Reels 9:16","salão se enchendo, a luz baixando"),
]
LIVRES=8
ini4=L4+2
for n in range(len(planejados)+LIVRES):
    r=ini4+n
    base=planejados[n] if n<len(planejados) else ("","","","","")
    for i,v in enumerate(base,1):
        if n<len(planejados):
            c=s4.cell(r,i,v); c.font=Font(F,size=10,bold=(i==1)); c.border=BORDA
            c.alignment=Alignment(vertical="center",wrap_text=True)
        else:
            entrada(s4.cell(r,i),None)
    entrada(s4.cell(r,6),None,DATA)
    for i in range(7,13): entrada(s4.cell(r,i),None,FMT_ENTRADA4[i])
    for i,(f_,fmt) in formulas_criativo(r).items(): calculo(s4.cell(r,i),f_,fmt)
    s4.row_dimensions[r].height=30
fim4=ini4+len(planejados)+LIVRES-1

dv_ang=DataValidation(type="list",formula1='"Destino,Ocasião,Processo,Produto,Vitrine"',allow_blank=True)
dv_fmt=DataValidation(type="list",formula1='"Reels 9:16,Foto 4:5,Carrossel 4:5,Vídeo 4:5"',allow_blank=True)
dv_camp=DataValidation(type="list",formula1='"QT_RES_WPP,QT_VIT_ALC,QT_GRP_WPP"',allow_blank=True)
for dv,col in [(dv_ang,"C"),(dv_fmt,"D"),(dv_camp,"B")]:
    s4.add_data_validation(dv); dv.add(f"{col}{ini4}:{col}{fim4}")

for col,w in zip("ABCDEFGHIJKLMNOPQ",[26,13,11,13,34,11,13,12,12,11,11,10,10,10,10,12,34]):
    s4.column_dimensions[col].width=w
s4.freeze_panes=f"B{L4+1}"


# ---------- checklist ----------
s5=wb.create_sheet("Checklist")
titulo(s5,"Checklist · semana 0 e lançamento","Marque sim na coluna Feito. O contador do topo soma sozinho.")
itens=[
 ("Semana 0",None,None),
 ("Criar o portfólio empresarial e conectar Página, Instagram e WhatsApp","Matheus","4.2"),
 ("Criar a conta de anúncios em Real e no fuso de São Paulo","Matheus","4.2"),
 ("Cadastrar pagamento e CNPJ","Matheus ou financeiro","4.2"),
 ("Dois administradores e autenticação em dois fatores para todos","Matheus","4.2"),
 ("Configurar saudação, ausência, respostas rápidas e etiquetas no WhatsApp Business","recepção","plano"),
 ("Definir o responsável pelo WhatsApp em cada turno e o tempo de resposta","gerente","4.4"),
 ("Coluna origem no livro de reservas e a pergunta: como conheceu a QT?","gerente","9.2"),
 ("Ticket médio sem os 13% levantado no Altec e lançado na Calculadora","Matheus","7.2"),
 ("Couverts de terça a quinta das 4 últimas semanas anotados","gerente","9.4"),
 ("Os 6 criativos separados ou gravados","Matheus e equipe","6"),
 ("Campanhas A e B montadas no Gerenciador, ainda sem publicar","Matheus","8"),
 ("Antes de publicar",None,None),
 ("Idade mínima de 18 em todos os conjuntos","Matheus","6.8"),
 ("Melhorias de criativo revisadas, as que alteram imagem e texto desligadas","Matheus","6.7"),
 ("Mensagem pré-preenchida diferente em cada campanha","Matheus","9.2"),
 ("Teste do caminho: na pré-visualização, tocar no botão e conferir se abre o WhatsApp da QT","Matheus","8.1"),
 ("Advertência de álcool nos anúncios em que o drink aparece","Matheus","6.8"),
 ("Texto dentro da zona segura do 9:16","Matheus","6.2"),
 ("Primeira semana",None,None),
 ("Publicado numa terça de manhã, com a equipe avisada","Matheus","8.1"),
 ("Três dias sem mexer: conferir só entrega, conversas e atendimento","Matheus","8.3"),
 ("Primeira linha da aba Semanal preenchida na segunda","Matheus","10.1"),
]
L5=5
cabecalho(s5,L5,["#","Tarefa","Quem","Apostila","Feito"])
r=L5+1; num=0
for t,quem,mod in itens:
    if quem is None:
        c=s5.cell(r,2,t); c.font=Font(F,size=10,bold=True,color=PRETO)
    else:
        num+=1
        s5.cell(r,1,num).font=Font(F,size=10)
        for i,v in [(2,t),(3,quem),(4,mod)]:
            c=s5.cell(r,i,v); c.font=Font(F,size=10); c.alignment=Alignment(wrap_text=True,vertical="center")
        entrada(s5.cell(r,5),None)
        for i in range(1,5): s5.cell(r,i).border=BORDA
    r+=1
fim5=r-1
dv_feito=DataValidation(type="list",formula1='"sim,não"',allow_blank=True)
s5.add_data_validation(dv_feito); dv_feito.add(f"E{L5+1}:E{fim5}")
s5["B3"]=f'="Feitos: "&COUNTIF(E{L5+1}:E{fim5},"sim")&" de {num}"'
s5["B3"].font=Font(F,size=11,bold=True,color=PRETO)
for col,w in zip("ABCDE",[5,74,20,10,9]): s5.column_dimensions[col].width=w


# ---------- como ler ----------
s6=wb.create_sheet("Como ler")
titulo(s6,"Como ler os números","As faixas de referência são ponto de partida. Depois de quatro semanas, o número que vale é o da própria QT.")
cabecalho(s6,4,["Métrica","O que é","Conta","Referência"])
gloss=[
 ("Alcance","pessoas diferentes que viram o anúncio","",""),
 ("Impressões","quantas vezes o anúncio apareceu","",""),
 ("Frequência","quantas vezes, em média, cada pessoa viu","impressões ÷ alcance","vitrine perto de 2 em 7 dias. Na reserva, acima de 3 com CTR caindo é cansaço"),
 ("CPM","custo de mil impressões, o preço do leilão","gasto ÷ impressões × 1.000","Brasil 2026, fontes de mercado: R$ 12 a R$ 26 na média, Reels e Stories de R$ 8 a R$ 20"),
 ("CTR do link","parte das impressões que virou toque no botão","cliques no link ÷ impressões","abaixo de 0,5%, o anúncio prende mas não convida"),
 ("Conversa iniciada","alguém mandou a primeira mensagem pelo anúncio","","é o resultado da campanha de reserva"),
 ("Custo por conversa","quanto custou cada conversa, sem tributo","gasto ÷ conversas","a meta está na Calculadora, seção 4"),
 ("Taxa de gancho","parte de quem viu que ficou 3 segundos","reproduções de 3 s ÷ impressões","abaixo de 20%, refazer o começo. De 25% a 30% para cima, funciona"),
 ("Retenção","de quem passou dos 3 segundos, quantos chegaram ao ThruPlay","ThruPlays ÷ reproduções de 3 s","compare entre os seus vídeos"),
 ("ThruPlay","vídeo assistido até o fim ou por 15 segundos","",""),
 ("Conversa que vira reserva","parte das conversas que fechou reserva","reservas ÷ conversas","abaixo de 15%, o problema está no atendimento ou na oferta"),
 ("Comparecimento","parte das reservas que apareceu","mesas que vieram ÷ reservas","confirmação na véspera sobe esse número"),
 ("Custo real por pessoa","o que a QT pagou, com tributo, por pessoa sentada","gasto real ÷ pessoas atendidas","a métrica que manda: compare com a meta e o empate da Calculadora"),
 ("ROAS real","faturamento por real de anúncio","faturamento ÷ gasto real","empate e meta na Calculadora"),
 ("% de mídia","o CMV do marketing","gasto real ÷ faturamento","empate e meta na Calculadora"),
 ("Fator do tributo","o que transforma o valor do Gerenciador em fatura","1 ÷ (1 − 12,15%)","1,1383. R$ 100 no Gerenciador são R$ 113,83 na fatura"),
]
r=5
for row in gloss:
    for i,v in enumerate(row,1):
        c=s6.cell(r,i,v); c.font=Font(F,size=10,bold=(i==1)); c.border=BORDA
        c.alignment=Alignment(wrap_text=True,vertical="top")
    r+=1

r+=1
secao(s6,r,"Regras de decisão · uma por segunda"); r+=1
cabecalho(s6,r,["Sinal","Leitura","O que fazer",""]); r+=1
regras=[
 ("Menos de 1.000 impressões ou menos de 3 dias","cedo demais","não mexer"),
 ("Anúncio gastou 2× a meta de custo por conversa sem nenhuma conversa","não funciona","pausar o anúncio"),
 ("Taxa de gancho abaixo de 20%","o começo não segura","refazer os 3 primeiros segundos"),
 ("Gancho bom e CTR do link abaixo de 0,5%","prende mas não convida","rever texto, oferta e botão"),
 ("Custo por conversa na meta e conversão em reserva abaixo de 15%","o problema não é o anúncio","rever tempo de resposta e roteiro do WhatsApp"),
 ("Frequência acima de 3 em 7 dias com CTR caindo","cansaço do criativo","entrar com criativo novo"),
 ("Custo real por pessoa abaixo da meta por 2 semanas seguidas","está dando certo","subir a verba em até 20%, a cada 3 ou 4 dias"),
 ("Custo real por pessoa acima do empate por 2 semanas seguidas","está perdendo dinheiro","baixar a verba ou pausar, e rever criativo e atendimento"),
]
for row in regras:
    for i,v in enumerate(row,1):
        c=s6.cell(r,i,v); c.font=Font(F,size=10); c.border=BORDA
        c.alignment=Alignment(wrap_text=True,vertical="top")
    r+=1

r+=1
secao(s6,r,"Não mexa · a câmara de fermentação"); r+=1
for t in ["Verba muda no máximo 20% de cada vez.",
          "No máximo uma mudança a cada 3 ou 4 dias, e anotada na coluna Mudança da aba Semanal.",
          "Nada de pausar e reativar toda hora, nem duplicar conjunto para reiniciar a sorte.",
          "Julgar pela semana, nunca pelo dia.",
          "No meio do aprendizado, entrar com no máximo dois anúncios novos de uma vez."]:
    c=s6.cell(r,1,"·"); c.font=Font(F,size=10)
    c=s6.cell(r,2,t); c.font=Font(F,size=10)
    r+=1
for col,w in zip("ABCD",[26,46,32,52]): s6.column_dimensions[col].width=w


# sem LibreOffice para gravar os valores, o Excel recalcula tudo ao abrir
wb.calculation.fullCalcOnLoad=True

saida=Path(__file__).resolve().parent/"meta-ads-qt.xlsx"
wb.save(saida)
print(f"salvo: {saida}")
